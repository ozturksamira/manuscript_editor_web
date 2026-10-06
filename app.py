import io
import json
import os
import re
import zipfile
from datetime import datetime
from functools import wraps
from html import escape
from html.parser import HTMLParser
from pathlib import Path

from lxml import etree
import docx
import PyPDF2
from flask import (
    Flask,
    jsonify,
    redirect,
    request,
    send_file,
    session,
    url_for,
)
from sqlalchemy import (
    Column,
    DateTime,
    ForeignKey,
    Integer,
    String,
    Text,
    UniqueConstraint,
    create_engine,
    inspect,
    select,
    text as sql_text,
)
from sqlalchemy.orm import declarative_base, relationship, sessionmaker
from werkzeug.security import check_password_hash, generate_password_hash


app = Flask(__name__)
PUBLIC_DIR = Path(__file__).resolve().parent / "public"


def frontend_page(filename):
    return send_file(PUBLIC_DIR / filename)

IS_RENDER = os.environ.get("RENDER", "").strip().lower() == "true"
SECRET_KEY = os.environ.get("SECRET_KEY", "").strip()
if IS_RENDER and not SECRET_KEY:
    raise RuntimeError("SECRET_KEY must be configured in Render.")
app.secret_key = SECRET_KEY or "dev-only-change-me"

DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
if IS_RENDER and not DATABASE_URL:
    raise RuntimeError(
        "DATABASE_URL must be configured in Render. Refusing ephemeral SQLite storage "
        "because it would lose accounts and manuscripts on redeploy."
    )
if DATABASE_URL:
    if DATABASE_URL.startswith("postgres://"):
        DATABASE_URL = "postgresql+psycopg://" + DATABASE_URL[len("postgres://"):]
    elif DATABASE_URL.startswith("postgresql://") and "+psycopg" not in DATABASE_URL:
        DATABASE_URL = DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)
else:
    DB_PATH = os.environ.get(
        "MM_DB_PATH",
        os.path.join(app.instance_path, "magnolia_margin.db"),
    )
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    DATABASE_URL = "sqlite:///" + DB_PATH

engine_kwargs = {"pool_pre_ping": True}
if DATABASE_URL.startswith("sqlite:"):
    engine_kwargs["connect_args"] = {"check_same_thread": False}

engine = create_engine(DATABASE_URL, **engine_kwargs)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
Base = declarative_base()

PROJECT_DOMAIN = os.environ.get("PROJECT_DOMAIN", "").strip().lower().rstrip(".")
PROJECT_SUBDOMAIN_PREFIX = os.environ.get("PROJECT_SUBDOMAIN_PREFIX", "").strip().lower()
if PROJECT_DOMAIN:
    app.config.update(
        SESSION_COOKIE_DOMAIN="." + PROJECT_DOMAIN,
        SESSION_COOKIE_SECURE=True,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
    )
else:
    app.config.update(
        SESSION_COOKIE_SECURE=IS_RENDER,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
    )


class User(Base):
    __tablename__ = "users"

    id = Column(Integer, primary_key=True)
    email = Column(String(320), nullable=False, unique=True, index=True)
    password_hash = Column(String(255), nullable=False)
    created_at = Column(DateTime, nullable=False, server_default=sql_text("CURRENT_TIMESTAMP"))

    projects = relationship(
        "Project",
        back_populates="user",
        cascade="all, delete-orphan",
        passive_deletes=True,
    )


class Project(Base):
    __tablename__ = "projects"
    __table_args__ = (
        UniqueConstraint("user_id", "slug", name="uq_project_user_slug"),
    )

    id = Column(Integer, primary_key=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    title = Column(String(160), nullable=False)
    slug = Column(String(180), nullable=True)
    source_filename = Column(String(255), nullable=False, default="")
    content = Column(Text, nullable=False, default="")
    analysis_json = Column(Text, nullable=False, default="")
    created_at = Column(DateTime, nullable=False, server_default=sql_text("CURRENT_TIMESTAMP"))
    updated_at = Column(
        DateTime,
        nullable=False,
        server_default=sql_text("CURRENT_TIMESTAMP"),
    )

    user = relationship("User", back_populates="projects")


ALLOWED_HTML_TAGS = {
    "p", "div", "h1", "h2", "h3", "h4", "h5", "h6",
    "strong", "b", "em", "i", "u", "s", "strike", "br",
    "blockquote", "ul", "ol", "li", "a", "hr", "span", "mark",
}
BLOCK_TAGS = {
    "p", "div", "h1", "h2", "h3", "h4", "h5", "h6",
    "blockquote", "li", "hr",
}
WORD_RE = re.compile(r"\b[A-Za-z](?:[A-Za-z'’-]*[A-Za-z])?\b")
WORD_COUNT_RE = re.compile(r"\b[A-Za-z](?:[A-Za-z'’-]*[A-Za-z])?\b")
SENTENCE_RE = re.compile(r"[^.!?\n]+(?:[.!?]+|$)")
PASSIVE_AUX_RE = re.compile(
    r"\b(?:am|is|are|was|were|be|been|being)\b"
    r"(?:\s+(?:not|never|already|still|just|quickly|slowly|suddenly|"
    r"clearly|completely|immediately|possibly|probably|often|usually|"
    r"nearly|almost|finally)){0,3}\s+"
    r"(?P<participle>[A-Za-z][A-Za-z'’-]*)\b",
    re.IGNORECASE,
)
GET_PASSIVE_RE = re.compile(
    r"\b(?:get|gets|got|getting|gotten)\b"
    r"(?:\s+(?:not|never|already|still|just|quickly|slowly|suddenly)){0,2}\s+"
    r"(?P<participle>[A-Za-z][A-Za-z'’-]*)\b",
    re.IGNORECASE,
)
IRREGULAR_PARTICIPLES = {
    "arisen", "awoken", "been", "begun", "bitten", "blown", "born", "bought",
    "bound", "broken", "brought", "built", "burnt", "caught", "chosen", "come",
    "cost", "cut", "dealt", "done", "drawn", "driven", "eaten", "fallen", "felt",
    "fought", "found", "flown", "forgiven", "forgotten", "frozen", "given", "gone",
    "grown", "heard", "held", "hidden", "hit", "hurt", "kept", "known", "laid",
    "led", "left", "lent", "let", "lost", "made", "meant", "met", "paid", "put",
    "read", "ridden", "run", "said", "seen", "sent", "set", "shaken", "shown",
    "shut", "sung", "sold", "spent", "spoken", "stood", "stolen", "stuck",
    "struck", "sworn", "swum", "taken", "taught", "torn", "told", "thought",
    "thrown", "understood", "woken", "won", "worn", "written",
}
AUXILIARY_WORDS = {
    "am", "is", "are", "was", "were", "be", "being", "been",
    "have", "has", "had", "having",
    "do", "does", "did",
    "can", "could", "may", "might", "must", "shall", "should",
    "will", "would",
    "ought", "need", "dare", "used",
}
PRONOUNS = {
    "i", "me", "my", "myself", "you", "your", "yourself", "yours",
    "he", "him", "his", "himself", "she", "her", "hers", "herself",
    "it", "its", "itself", "we", "us", "our", "ourselves",
    "they", "them", "their", "theirs", "themselves",
    "who", "whom", "whose", "which", "that", "this", "these", "those",
}
PREPOSITIONS = {
    "about", "above", "across", "after", "against", "along", "among", "around",
    "at", "before", "behind", "below", "beneath", "beside", "between", "beyond",
    "by", "down", "during", "except", "for", "from", "in", "inside", "into",
    "like", "near", "of", "off", "on", "onto", "out", "outside", "over",
    "past", "through", "throughout", "to", "toward", "under", "underneath",
    "until", "up", "upon", "with", "within", "without",
}
CONJUNCTIONS = {
    "and", "but", "or", "nor", "for", "yet", "so", "although", "because",
    "since", "unless", "until", "while", "whereas",
}
CONTRACTION_STOP_WORDS = {
    "i'm", "i’ve", "i'll", "i’d",
    "you're", "you’ve", "you'll", "you’d",
    "he's", "he’ll", "he’d",
    "she's", "she’ll", "she’d",
    "it's", "it’ll", "it’d",
    "we're", "we’ve", "we'll", "we’d",
    "they're", "they’ve", "they'll", "they’d",
    "who's", "who’ll", "who’d",
    "that's", "that’ll", "that’d",
    "can't", "couldn't", "won't", "wouldn't", "shouldn't", "shan't",
    "isn't", "aren't", "wasn't", "weren't", "ain't",
    "haven't", "hasn't", "hadn't",
    "don't", "doesn't", "didn't",
    "can't", "could've", "couldn’t", "would've", "wouldn’t",
    "should've", "shouldn’t", "might've", "mightn’t", "must've", "mustn’t",
    "needn't", "daren't", "oughtn't", "usedn't",
    "cannot",
}
STOP_WORDS = PRONOUNS | PREPOSITIONS | CONJUNCTIONS | AUXILIARY_WORDS | CONTRACTION_STOP_WORDS | {
    "a", "an", "the",
}
EVENT_VERBS = {
    "arrive", "arrived", "attack", "attacked", "avoid", "avoided", "break", "broke",
    "burst", "build", "built", "call", "called", "catch", "caught", "change", "changed",
    "chase", "chased", "choose", "chose", "close", "closed", "crash", "crashed",
    "cry", "cried", "cut", "cutting", "die", "died", "discover", "discovered",
    "drag", "dragged", "drive", "drove", "escape", "escaped", "enter", "entered",
    "explode", "exploded", "fall", "fell", "fight", "fought", "find", "found",
    "flee", "fled", "grab", "grabbed", "hit", "hold", "held", "jump", "jumped",
    "kick", "kicked", "kill", "killed", "leave", "left", "lose", "lost", "open", "opened",
    "pull", "pulled", "push", "pushed", "reach", "reached", "raise", "raised",
    "react", "reacted", "run", "ran", "rush", "rushed", "scream", "screamed",
    "search", "searched", "see", "saw", "send", "sent", "shake", "shook", "shoot",
    "shot", "shout", "shouted", "slam", "slammed", "slip", "slipped", "smash",
    "smashed", "sprint", "sprinted", "stand", "stood", "start", "started",
    "stop", "stopped", "strike", "struck", "survive", "survived", "take", "took",
    "throw", "threw", "turn", "turned", "wake", "woke", "walk", "walked",
    "warn", "warned", "watch", "watched", "whisper", "whispered", "write", "wrote",
}
DESCRIPTION_RE = re.compile(
    r"\b(?:very|quite|rather|really|extremely|beautifully|slowly|quickly|"
    r"carefully|suddenly|silently|quietly|loudly|softly|deeply|"
    r"[A-Za-z-]*(?:ly|ful|ous|ive|less|ish|ical))\b",
    re.IGNORECASE,
)



def slugify(value, fallback="project"):
    value = (value or "").strip().lower()
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value[:150] or fallback


def unique_slug(db, user_id, title, project_id=None):
    base = slugify(title)
    candidate = base
    n = 2
    while True:
        query = select(Project.id).where(
            Project.user_id == user_id,
            Project.slug == candidate,
        )
        if project_id:
            query = query.where(Project.id != project_id)
        if db.execute(query).first() is None:
            return candidate
        candidate = f"{base}-{n}"
        n += 1


def project_host_slug():
    if not PROJECT_DOMAIN:
        return None
    host = request.host.split(":", 1)[0].lower()
    suffix = "." + PROJECT_DOMAIN
    if not host.endswith(suffix):
        return None
    prefix = host[: -len(suffix)]
    if PROJECT_SUBDOMAIN_PREFIX:
        expected = PROJECT_SUBDOMAIN_PREFIX + "-"
        if not prefix.startswith(expected):
            return None
        prefix = prefix[len(expected):]
    if re.fullmatch(r"[a-z0-9-]{1,150}", prefix):
        return prefix
    return None


def project_link(project):
    if PROJECT_DOMAIN:
        prefix = (PROJECT_SUBDOMAIN_PREFIX + "-") if PROJECT_SUBDOMAIN_PREFIX else ""
        return f"https://{prefix}{project.slug}.{PROJECT_DOMAIN}/"
    return url_for("project_editor", slug=project.slug)


class Sanitizer(HTMLParser):
    SAFE_STYLE_PROPERTIES = {
        "color", "background-color", "font-size", "text-align",
        "font-family", "letter-spacing", "margin-left", "margin-right",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def safe_style(self, style):
        safe = []
        for declaration in (style or "").split(";"):
            if ":" not in declaration:
                continue
            name, value = declaration.split(":", 1)
            name = name.strip().lower()
            value = value.strip()
            if name not in self.SAFE_STYLE_PROPERTIES:
                continue
            if re.fullmatch(r"(?:#[0-9a-f]{3,8}|[a-zA-Z ]+|[0-9.]+(?:px|pt|em|rem|%)?)", value):
                safe.append(f"{name}:{value}")
        return ";".join(safe)

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag not in ALLOWED_HTML_TAGS:
            return

        if tag == "a":
            href = ""
            for key, value in attrs:
                if key.lower() == "href" and value:
                    href = value.strip()
            if href.startswith(("https://", "http://", "mailto:")):
                self.parts.append(
                    f'<a href="{escape(href, quote=True)}" target="_blank" rel="noopener">'
                )
            else:
                self.parts.append("<a>")
            return

        if tag == "span":
            styles = ""
            for key, value in attrs:
                if key.lower() == "style":
                    styles = self.safe_style(value)
            if styles:
                self.parts.append(f'<span style="{escape(styles, quote=True)}">')
            else:
                self.parts.append("<span>")
            return

        if tag in {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6"}:
            styles = ""
            for key, value in attrs:
                if key.lower() == "style":
                    styles = self.safe_style(value)
            self.parts.append(f"<{tag}{(' style=' + chr(34) + escape(styles, quote=True) + chr(34)) if styles else ''}>")
            return

        self.parts.append(f"<{tag}>")

    def handle_startendtag(self, tag, attrs):
        if tag.lower() == "br":
            self.parts.append("<br>")
        elif tag.lower() == "hr":
            self.parts.append("<hr>")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ALLOWED_HTML_TAGS and tag not in {"br", "hr"}:
            self.parts.append(f"</{tag}>")

    def handle_data(self, data):
        self.parts.append(escape(data))


class TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def _newline(self):
        if self.parts and not self.parts[-1].endswith("\n"):
            self.parts.append("\n")

    def handle_starttag(self, tag, attrs):
        if tag.lower() in BLOCK_TAGS:
            self._newline()

    def handle_endtag(self, tag):
        if tag.lower() in BLOCK_TAGS:
            self._newline()

    def handle_data(self, data):
        self.parts.append(data)


def sanitize_html(value):
    parser = Sanitizer()
    parser.feed(value or "")
    parser.close()
    return "".join(parser.parts).strip()


def html_to_text(content):
    parser = TextExtractor()
    parser.feed(content or "")
    parser.close()
    value = "".join(parser.parts).replace("\xa0", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r" *\n *", "\n", value)
    return value.strip()


def text_to_html(value):
    paragraphs = re.split(r"\n\s*\n", (value or "").strip())
    if paragraphs == [""]:
        return "<p></p>"
    return "".join(
        f"<p>{escape(paragraph).replace(chr(10), '<br>')}</p>"
        for paragraph in paragraphs
    )


def _xml_hex_color(node, xpath):
    value = node.find(xpath, namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
    return value.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") if value is not None else None


def _run_html(run):
    text = "".join(run.itertext())
    if not text:
        return ""
    output = escape(text).replace("\n", "<br>")

    rpr = run.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
    styles = []
    if rpr is not None:
        color = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color")
        if color is not None and color.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"):
            value = color.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
            if re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                styles.append(f"color:#{value}")

        size = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz")
        if size is not None and size.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"):
            try:
                points = int(size.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")) / 2
                styles.append(f"font-size:{points:g}pt")
            except ValueError:
                pass

        highlight = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}highlight")
        if highlight is not None and highlight.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"):
            highlight_value = highlight.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
            palette = {
                "yellow": "#fff3a3", "green": "#d8ead8", "cyan": "#d7f0f2",
                "magenta": "#f0d8e9", "blue": "#dbe8f7", "red": "#f4d6d1",
            }
            if highlight_value in palette:
                styles.append(f"background-color:{palette[highlight_value]}")

        if rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}b") is not None:
            output = f"<strong>{output}</strong>"
        if rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}i") is not None:
            output = f"<em>{output}</em>"
        if rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}u") is not None:
            output = f"<u>{output}</u>"
        if rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}strike") is not None:
            output = f"<s>{output}</s>"

    if styles:
        output = f'<span style="{";".join(styles)}">{output}</span>'
    return output


def docx_to_html(raw_bytes):
    # Read the DOCX package directly. In addition to being faster for large
    # manuscripts, this makes heading detection independent of Word's style
    # naming conventions: outline levels and style definitions are both checked.
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    qn = lambda local: "{%s}%s" % (ns["w"], local)

    with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
        document_root = etree.fromstring(archive.read("word/document.xml"))
        style_map = {}

        if "word/styles.xml" in archive.namelist():
            styles_root = etree.fromstring(archive.read("word/styles.xml"))
            for style in styles_root.xpath(".//w:style[@w:type='paragraph']", namespaces=ns):
                style_id = style.get(qn("styleId"), "")
                name_node = style.find("w:name", ns)
                style_name = name_node.get(qn("val"), "") if name_node is not None else ""
                outline_node = style.find(".//w:outlineLvl", ns)
                outline_level = outline_node.get(qn("val")) if outline_node is not None else None

                level = None
                match = re.search(r"(?:heading|titre)\s*([1-6])", style_name, re.IGNORECASE)
                if not match:
                    match = re.search(r"(?:heading|titre)([1-6])", style_id, re.IGNORECASE)
                if match:
                    level = int(match.group(1))
                elif outline_level is not None and outline_level.isdigit():
                    level = min(6, int(outline_level) + 1)

                if style_id:
                    style_map[style_id] = {"name": style_name, "level": level}

    blocks = []
    body = document_root.find("w:body", ns)
    paragraphs = body.findall("w:p", ns) if body is not None else []

    for paragraph in paragraphs:
        ppr = paragraph.find("w:pPr", ns)
        style_id = ""
        if ppr is not None:
            pstyle = ppr.find("w:pStyle", ns)
            if pstyle is not None:
                style_id = pstyle.get(qn("val"), "") or ""

        style_info = style_map.get(style_id, {})
        style_name = style_info.get("name", "") or ""
        heading_level = style_info.get("level")

        if ppr is not None and heading_level is None:
            outline_node = ppr.find("w:outlineLvl", ns)
            if outline_node is not None and outline_node.get(qn("val"), "").isdigit():
                heading_level = min(6, int(outline_node.get(qn("val"))) + 1)

        heading_match = re.search(
            r"(?:heading|titre)\s*([1-6])",
            style_name,
            flags=re.IGNORECASE,
        )
        if heading_match:
            heading_level = int(heading_match.group(1))

        run_html = []
        for child in paragraph:
            tag = etree.QName(child).localname
            if tag == "r":
                run_html.append(_run_html(child))
            elif tag == "hyperlink":
                for run in child.xpath(".//w:r", namespaces=ns):
                    run_html.append(_run_html(run))

        paragraph_body = "".join(run_html)
        if not paragraph_body:
            paragraph_body = "<br>"

        paragraph_styles = []
        if ppr is not None:
            jc = ppr.find("w:jc", ns)
            if jc is not None:
                align = jc.get(qn("val"), "")
                if align in {"left", "center", "right", "both", "justify"}:
                    paragraph_styles.append(
                        "text-align:" + ("justify" if align in {"both", "justify"} else align)
                    )

            shd = ppr.find("w:shd", ns)
            fill = shd.get(qn("fill"), "") if shd is not None else ""
            if re.fullmatch(r"[0-9A-Fa-f]{6}", fill or ""):
                paragraph_styles.append(f"background-color:#{fill}")

            borders = ppr.find("w:pBdr", ns)
            if borders is not None:
                separator_found = False
                for side_name in ("bottom", "top"):
                    side = borders.find(f"w:{side_name}", ns)
                    if side is not None and side.get(qn("val"), "") not in {"", "nil"}:
                        separator_found = True
                        break
                if separator_found:
                    blocks.append("<hr>")

        style_attr = (
            ' style="' + escape(";".join(paragraph_styles), quote=True) + '"'
            if paragraph_styles else ""
        )

        if heading_level:
            tag = f"h{heading_level}"
        elif "quote" in style_name.lower():
            tag = "blockquote"
        elif style_name.lower() in {"title", "subtitle"}:
            tag = "h1" if style_name.lower() == "title" else "h2"
        else:
            tag = "p"

        blocks.append(f"<{tag}{style_attr}>{paragraph_body}</{tag}>")

    return "".join(blocks)


def pdf_to_html(raw_bytes):
    reader = PyPDF2.PdfReader(io.BytesIO(raw_bytes))
    pages = []
    for page in reader.pages:
        value = page.extract_text() or ""
        if value.strip():
            pages.append(value.strip())
    return text_to_html("\n\n".join(pages))


def parse_uploaded_file(file_storage):
    filename = (file_storage.filename or "").strip()
    lower = filename.lower()
    raw = file_storage.read()

    if lower.endswith(".txt"):
        return filename, text_to_html(raw.decode("utf-8-sig", errors="replace"))
    if lower.endswith(".docx"):
        return filename, docx_to_html(raw)
    if lower.endswith(".pdf"):
        return filename, pdf_to_html(raw)
    raise ValueError("Unsupported file type. Please use .txt, .docx, or .pdf.")


def sentence_spans(value):
    output = []
    for match in SENTENCE_RE.finditer(value or ""):
        snippet = match.group(0).strip()
        if not snippet:
            continue
        leading = len(match.group(0)) - len(match.group(0).lstrip())
        trailing = len(match.group(0).rstrip())
        output.append((match.start() + leading, match.start() + trailing, snippet))
    return output


def _sentence_profile(snippet):
    words = [
        match.group(0).lower().replace("’", "'").strip("'")
        for match in WORD_RE.finditer(snippet)
    ]
    content_words = [word for word in words if word not in STOP_WORDS]
    unique_content = len(set(content_words))
    word_count = len(words)
    content_count = len(content_words)
    event_count = sum(1 for word in words if word in EVENT_VERBS)
    description_count = len(DESCRIPTION_RE.findall(snippet))
    dialogue = 1 if re.search(r'["“”]', snippet) else 0
    questions = snippet.count("?")
    exclamations = snippet.count("!")
    return {
        "word_count": word_count,
        "content_word_count": content_count,
        "content_variety": (unique_content / content_count) if content_count else 1.0,
        "event_count": event_count,
        "event_density": (event_count / word_count) if word_count else 0.0,
        "description_count": description_count,
        "description_density": (description_count / word_count) if word_count else 0.0,
        "dialogue": dialogue,
        "questions": questions,
        "exclamations": exclamations,
    }

def _is_participle(word):
    word = word.lower().strip("'")
    return (
        word in IRREGULAR_PARTICIPLES
        or word.endswith("ed")
        or word.endswith("en")
        or word.endswith("wn")
        or word.endswith("t")
    )


def _passive_findings(snippet, start):
    findings = []
    seen_spans = set()

    for pattern in (PASSIVE_AUX_RE, GET_PASSIVE_RE):
        for match in pattern.finditer(snippet):
            participle = match.group("participle")
            if not _is_participle(participle):
                continue
            span = (match.start(), match.end())
            if span in seen_spans:
                continue
            seen_spans.add(span)

            trailing = snippet[match.end():match.end() + 90]
            agent_match = re.search(
                r"\bby\s+(?:the|a|an)?\s*[A-Za-z][A-Za-z'’-]*\b",
                trailing,
                re.IGNORECASE,
            )
            has_agent = bool(agent_match)
            reason = (
                "High-confidence passive voice — explicit agent"
                if has_agent
                else "Possible passive voice — action recipient is in subject position"
            )
            findings.append({
                "sentence": snippet[:500],
                "start": start,
                "end": start + len(snippet),
                "reason": reason,
                "confidence": "high" if has_agent else "medium",
            })
    return findings


def analyse_text(value):
    value = value or ""
    if not value.strip():
        return {
            "word_count": 0,
            "words": [],
            "passive": [],
            "pacing": [],
            "flagged_count": 0,
        }

    sentence_ranges = sentence_spans(value)
    word_count = sum(1 for _ in WORD_COUNT_RE.finditer(value))
    counts = {}

    for match in WORD_RE.finditer(value):
        word = match.group(0).lower().replace("’", "'").strip("'")
        if len(word) < 2 or word in STOP_WORDS:
            continue
        counts[word] = counts.get(word, 0) + 1

    frequent = sorted(
        ((word, count) for word, count in counts.items() if count > 3),
        key=lambda pair: (-pair[1], pair[0]),
    )[:100]
    frequent_set = {word for word, _ in frequent}

    contexts = {}
    for start, end, snippet in sentence_ranges:
        if frequent_set and len(contexts) >= len(frequent_set):
            break
        for match in WORD_RE.finditer(snippet):
            word = match.group(0).lower().replace("’", "'").strip("'")
            if word in frequent_set:
                bucket = contexts.setdefault(word, [])
                if len(bucket) < 3 and snippet[:300] not in bucket:
                    bucket.append(snippet[:300])

    words = [
        {"word": word, "count": count, "contexts": contexts.get(word, [])}
        for word, count in frequent
    ]

    profiles = []
    passive = []
    for start, end, snippet in sentence_ranges:
        profile = _sentence_profile(snippet)
        profile.update({"start": start, "end": end, "sentence": snippet})
        profiles.append(profile)
        passive.extend(_passive_findings(snippet, start))
        if len(passive) >= 120:
            passive = passive[:120]
            # Continue building sentence profiles for pacing.

    pacing = []

    def window_text(window):
        return " ".join(item["sentence"].strip() for item in window)

    # Too fast / event compression: several short sentences with concentrated
    # action, suggesting that major beats may be arriving before the reader can
    # process the change, reaction, or consequence.
    for i in range(max(0, len(profiles) - 2)):
        window = profiles[i:i + 3]
        total_words = sum(item["word_count"] for item in window)
        avg_words = total_words / 3
        event_total = sum(item["event_count"] for item in window)
        event_density = event_total / total_words if total_words else 0.0
        short_sentences = sum(item["word_count"] <= 15 for item in window)
        if (
            short_sentences >= 2
            and event_total >= 2
            and avg_words <= 17
            and event_density >= 0.055
        ):
            pacing.append({
                "sentence": window_text(window)[:650],
                "start": window[0]["start"],
                "end": window[-1]["end"],
                "reason": "Too fast — possible event compression",
                "word_count": total_words,
                "confidence": "medium",
            })
        if len(pacing) >= 80:
            break

    # Too slow / dragging: several long sentences with little event movement,
    # heavy description, repetitive content vocabulary, or dialogue that keeps
    # the plot in place without a change in stakes.
    if len(pacing) < 80:
        for i in range(max(0, len(profiles) - 3)):
            window = profiles[i:i + 4]
            total_words = sum(item["word_count"] for item in window)
            avg_words = total_words / 4
            event_total = sum(item["event_count"] for item in window)
            event_density = event_total / total_words if total_words else 0.0
            description_density = sum(item["description_count"] for item in window) / total_words if total_words else 0.0
            dialogue_rate = sum(item["dialogue"] for item in window) / 4
            variety = sum(item["content_variety"] for item in window) / 4
            drag_signal = (
                description_density >= 0.06
                or dialogue_rate >= 0.5
                or variety <= 0.62
            )
            if (
                avg_words >= 12
                and event_total <= 2
                and event_density <= 0.04
                and drag_signal
            ):
                pacing.append({
                    "sentence": window_text(window)[:700],
                    "start": window[0]["start"],
                    "end": window[-1]["end"],
                    "reason": "Too slow — possible dragging passage",
                    "word_count": total_words,
                    "confidence": "medium",
                })
            if len(pacing) >= 80:
                break

    # Flat dynamics: a run of medium-length sentences with almost identical
    # rhythm and event activity, indicating a sustained single-gear passage.
    if len(pacing) < 100 and len(profiles) >= 8:
        for i in range(len(profiles) - 7):
            window = profiles[i:i + 8]
            lengths = [item["word_count"] for item in window]
            events = [item["event_count"] for item in window]
            mean_len = sum(lengths) / 8
            variance = sum((x - mean_len) ** 2 for x in lengths) / 8
            stdev = variance ** 0.5
            event_range = max(events) - min(events)
            rhythm_scores = [
                (item["word_count"] / 8.0)
                + (item["event_count"] * 2.0)
                + (item["dialogue"] * 0.6)
                + ((item["questions"] + item["exclamations"]) * 0.5)
                for item in window
            ]
            rhythm_mean = sum(rhythm_scores) / 8
            rhythm_variance = sum((score - rhythm_mean) ** 2 for score in rhythm_scores) / 8
            rhythm_stdev = rhythm_variance ** 0.5
            if (
                5 <= mean_len <= 28
                and stdev <= 4.5
                and event_range <= 1
                and rhythm_stdev <= 1.7
            ):
                pacing.append({
                    "sentence": window_text(window)[:750],
                    "start": window[0]["start"],
                    "end": window[-1]["end"],
                    "reason": "Flat dynamics — sustained single-gear rhythm",
                    "word_count": sum(lengths),
                    "confidence": "medium",
                })
            if len(pacing) >= 100:
                break

    # Deduplicate overlapping findings of the same type so one slow or flat
    # stretch produces one useful finding rather than a card for every shift.
    deduped = []
    last_end_by_reason = {}
    for item in pacing:
        reason = item["reason"]
        previous_end = last_end_by_reason.get(reason, -1)
        if item["start"] <= previous_end:
            continue
        deduped.append(item)
        last_end_by_reason[reason] = item["end"]

    pacing = deduped[:100]
    passive = passive[:120]

    return {
        "word_count": word_count,
        "words": words,
        "passive": passive,
        "pacing": pacing,
        "flagged_count": len(words) + len(passive) + len(pacing),
    }

def get_db():
    return SessionLocal()


def migrate_schema():
    Base.metadata.create_all(engine)

    inspector = inspect(engine)
    tables = set(inspector.get_table_names())

    # The previous version used SQLite/SQL statements directly. Add the new
    # slug column to that schema if an existing local database is reused.
    if "projects" in tables:
        columns = {column["name"] for column in inspector.get_columns("projects")}
        if "slug" not in columns:
            with engine.begin() as connection:
                connection.execute(sql_text("ALTER TABLE projects ADD COLUMN slug VARCHAR(180)"))

    db = get_db()
    try:
        users = db.execute(select(User)).scalars().all()
        for user in users:
            legacy_projects = db.execute(
                select(Project)
                .where(Project.user_id == user.id)
                .order_by(Project.id)
            ).scalars().all()

            if not legacy_projects and "documents" in tables:
                # Best-effort migration from the previous single-document schema.
                rows = db.execute(
                    sql_text(
                        "SELECT content FROM documents WHERE user_id = :user_id"
                    ),
                    {"user_id": user.id},
                ).fetchall()
                legacy_content = rows[0][0] if rows else ""
                project = Project(
                    user_id=user.id,
                    title="My manuscript",
                    slug=None,
                    content=sanitize_html(legacy_content),
                )
                db.add(project)
                db.flush()
                legacy_projects = [project]

            for project in legacy_projects:
                if not project.slug:
                    project.slug = unique_slug(db, user.id, project.title, project.id)

        db.commit()
    finally:
        db.close()


def current_user():
    user_id = session.get("user_id")
    if not user_id:
        return None
    db = get_db()
    try:
        return db.get(User, int(user_id))
    finally:
        db.close()


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not current_user():
            return redirect(url_for("login_page"))
        return view(*args, **kwargs)
    return wrapped


def api_login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not current_user():
            return jsonify({"error": "Authentication required."}), 401
        return view(*args, **kwargs)
    return wrapped


def project_for_user(db, project_id, user_id):
    return db.execute(
        select(Project).where(
            Project.id == project_id,
            Project.user_id == user_id,
        )
    ).scalar_one_or_none()


def project_for_slug(db, slug, user_id):
    return db.execute(
        select(Project).where(
            Project.slug == slug,
            Project.user_id == user_id,
        )
    ).scalar_one_or_none()


def project_payload(project, include_content=False):
    payload = {
        "id": project.id,
        "title": project.title,
        "slug": project.slug,
        "source_filename": project.source_filename or "",
        "created_at": project.created_at.isoformat() if project.created_at else None,
        "updated_at": project.updated_at.isoformat() if project.updated_at else None,
        "url": project_link(project),
    }
    if include_content:
        payload["content"] = project.content or ""
        payload["analysis"] = json.loads(project.analysis_json or "{}")
    return payload


@app.route("/")
def index():
    host_slug = project_host_slug()
    if host_slug:
        user = current_user()
        if not user:
            return redirect(url_for("login_page"))
        db = get_db()
        try:
            project = project_for_slug(db, host_slug, user.id)
            if not project:
                return frontend_page("not_found.html"), 404
            return frontend_page("editor.html")
        finally:
            db.close()
    return frontend_page("index.html")


@app.route("/login/")
@app.route("/login")
def login_page():
    if current_user():
        return redirect(url_for("projects"))
    return frontend_page("login.html")


@app.route("/projects/")
@app.route("/projects")
@login_required
def projects():
    return frontend_page("projects.html")


@app.route("/editor/")
@app.route("/editor")
@login_required
def editor():
    db = get_db()
    try:
        first_project = db.execute(
            select(Project).where(Project.user_id == current_user().id).order_by(Project.updated_at.desc(), Project.id.desc())
        ).scalars().first()
        if not first_project:
            first_project = Project(
                user_id=current_user().id,
                title="Untitled manuscript",
                slug=None,
                content="",
            )
            db.add(first_project)
            db.flush()
            first_project.slug = unique_slug(db, current_user().id, first_project.title, first_project.id)
            db.commit()
        return redirect(project_link(first_project))
    finally:
        db.close()


@app.route("/project/<slug>/")
@app.route("/project/<slug>")
@login_required
def project_editor(slug):
    db = get_db()
    try:
        project = project_for_slug(db, slug, current_user().id)
        if not project:
            return frontend_page("not_found.html"), 404
        return frontend_page("editor.html")
    finally:
        db.close()


@app.route("/auth", methods=["POST"])
def auth():
    data = request.get_json(silent=True) or request.form
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""
    action = data.get("action", "login")

    if not email or not password:
        return jsonify({"error": "Email and password are required."}), 400
    if len(password) < 8:
        return jsonify({"error": "Your password must be at least 8 characters."}), 400

    db = get_db()
    try:
        user = db.execute(select(User).where(User.email == email)).scalar_one_or_none()

        if action == "register":
            if user:
                return jsonify({"error": "An account with that email already exists. Please sign in."}), 409

            user = User(
                email=email,
                password_hash=generate_password_hash(password),
            )
            db.add(user)
            db.flush()

            project = Project(
                user_id=user.id,
                title="Untitled manuscript",
                slug=unique_slug(db, user.id, "Untitled manuscript"),
                content="",
            )
            db.add(project)
            db.commit()

            session.clear()
            session["user_id"] = user.id
            session["email"] = user.email
            return jsonify({"status": "success", "message": "Account created successfully."})

        if action == "login":
            if user and check_password_hash(user.password_hash, password):
                session.clear()
                session["user_id"] = user.id
                session["email"] = user.email
                return jsonify({"status": "success", "message": "Logged in successfully."})
            return jsonify({"error": "Invalid email or password."}), 401

        return jsonify({"error": "Unknown authentication action."}), 400
    finally:
        db.close()


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"status": "success"})


@app.route("/api/account/delete", methods=["POST"])
@api_login_required
def delete_account():
    user = current_user()
    data = request.get_json(silent=True) or {}
    password = data.get("password") or ""
    confirmation = (data.get("confirmation") or "").strip()

    if confirmation != "DELETE":
        return jsonify({"error": "Type DELETE to confirm permanent account removal."}), 400
    if not check_password_hash(user.password_hash, password):
        return jsonify({"error": "Your password is incorrect."}), 401

    db = get_db()
    try:
        db.delete(db.get(User, user.id))
        db.commit()
    finally:
        db.close()

    session.clear()
    return jsonify({"status": "deleted"})


@app.route("/api/projects", methods=["GET"])
@api_login_required
def list_projects():
    user = current_user()
    db = get_db()
    try:
        projects = db.execute(
            select(Project)
            .where(Project.user_id == user.id)
            .order_by(Project.updated_at.desc(), Project.id.desc())
        ).scalars().all()
        return jsonify({"projects": [project_payload(project) for project in projects]})
    finally:
        db.close()


@app.route("/api/projects", methods=["POST"])
@api_login_required
def create_project():
    user = current_user()
    data = request.get_json(silent=True) or {}
    title = (data.get("title") or "Untitled manuscript").strip()[:160] or "Untitled manuscript"

    db = get_db()
    try:
        project = Project(
            user_id=user.id,
            title=title,
            slug=unique_slug(db, user.id, title),
            content="",
        )
        db.add(project)
        db.commit()
        db.refresh(project)
        return jsonify({"project": project_payload(project, include_content=True)}), 201
    finally:
        db.close()


@app.route("/api/projects/<int:project_id>", methods=["GET"])
@api_login_required
def get_project(project_id):
    user = current_user()
    db = get_db()
    try:
        project = project_for_user(db, project_id, user.id)
        if not project:
            return jsonify({"error": "Project not found."}), 404
        return jsonify({"project": project_payload(project, include_content=True)})
    finally:
        db.close()


@app.route("/api/projects/by-slug/<slug>", methods=["GET"])
@api_login_required
def get_project_by_slug(slug):
    user = current_user()
    db = get_db()
    try:
        project = project_for_slug(db, slug, user.id)
        if not project:
            return jsonify({"error": "Project not found."}), 404
        return jsonify({"project": project_payload(project, include_content=True)})
    finally:
        db.close()


@app.route("/api/projects/<int:project_id>", methods=["POST"])
@api_login_required
def save_project(project_id):
    user = current_user()
    db = get_db()
    try:
        project = project_for_user(db, project_id, user.id)
        if not project:
            return jsonify({"error": "Project not found."}), 404

        data = request.get_json(silent=True) or {}
        title = (data.get("title") or project.title).strip()[:160] or "Untitled manuscript"
        content = sanitize_html(data.get("content", project.content or ""))
        analysis = data.get("analysis")
        if isinstance(analysis, dict):
            project.analysis_json = json.dumps(analysis, ensure_ascii=False)

        if title != project.title:
            project.title = title
            project.slug = unique_slug(db, user.id, title, project.id)

        project.content = content
        project.updated_at = datetime.utcnow()
        db.commit()

        return jsonify({
            "status": "success",
            "updated_at": project.updated_at.isoformat(),
            "slug": project.slug,
            "url": project_link(project),
        })
    finally:
        db.close()


@app.route("/api/projects/<int:project_id>/analysis", methods=["POST"])
@api_login_required
def save_analysis(project_id):
    user = current_user()
    db = get_db()
    try:
        project = project_for_user(db, project_id, user.id)
        if not project:
            return jsonify({"error": "Project not found."}), 404
        data = request.get_json(silent=True) or {}
        analysis = data.get("analysis") or {}
        if not isinstance(analysis, dict):
            return jsonify({"error": "Invalid analysis payload."}), 400
        project.analysis_json = json.dumps(analysis, ensure_ascii=False)
        project.updated_at = datetime.utcnow()
        db.commit()
        return jsonify({"status": "success"})
    finally:
        db.close()


@app.route("/api/projects/<int:project_id>", methods=["DELETE"])
@api_login_required
def delete_project(project_id):
    user = current_user()
    db = get_db()
    try:
        project = project_for_user(db, project_id, user.id)
        if not project:
            return jsonify({"error": "Project not found."}), 404

        count = db.execute(
            select(Project.id).where(Project.user_id == user.id)
        ).all()
        if len(count) <= 1:
            return jsonify({"error": "Keep at least one project in your writing room."}), 400

        db.delete(project)
        db.commit()
        return jsonify({"status": "success"})
    finally:
        db.close()


@app.route("/api/projects/import", methods=["POST"])
@api_login_required
def import_project():
    user = current_user()
    filename = ""
    content = ""
    uploaded_file = request.files.get("file")
    if uploaded_file and uploaded_file.filename:
        try:
            filename, content = parse_uploaded_file(uploaded_file)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception as exc:
            return jsonify({"error": f"Failed to read the manuscript: {exc}"}), 500
    else:
        data = request.get_json(silent=True) or {}
        filename = str(data.get("filename") or "").strip()[:255]
        content = str(data.get("content") or "")
        if not filename or not content:
            return jsonify({"error": "Choose a manuscript file first."}), 400
        if not re.search(r"\.(txt|docx|pdf)$", filename, re.IGNORECASE):
            return jsonify({"error": "Unsupported file type. Please use .txt, .docx, or .pdf."}), 400

    content = sanitize_html(content)
    title = str((request.get_json(silent=True) or {}).get("title") or "").strip()[:160] if not uploaded_file else ""
    if not title:
        title = Path(filename).stem.strip()[:160] or "Imported manuscript"

    db = get_db()
    try:
        project = Project(
            user_id=user.id,
            title=title,
            slug=unique_slug(db, user.id, title),
            source_filename=filename,
            content=content,
        )
        db.add(project)
        db.commit()
        db.refresh(project)
        return jsonify({"project": project_payload(project, include_content=True)}), 201
    finally:
        db.close()


@app.route("/api/projects/<int:project_id>/analyze", methods=["POST"])
@api_login_required
def analyse_project(project_id):
    user = current_user()
    db = get_db()
    try:
        project = project_for_user(db, project_id, user.id)
        if not project:
            return jsonify({"error": "Project not found."}), 404
        data = request.get_json(silent=True) or {}
        manuscript_text = data.get("text")
        if manuscript_text is None:
            manuscript_text = html_to_text(project.content)
        analysis = analyse_text(manuscript_text)
        project.analysis_json = json.dumps(analysis, ensure_ascii=False)
        project.updated_at = datetime.utcnow()
        db.commit()
        return jsonify(analysis)
    finally:
        db.close()


def html_to_docx(content):
    class DocumentBuilder(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.document = docx.Document()
            self.paragraph = None
            self.bold = False
            self.italic = False
            self.underline = False
            self.strike = False

        def ensure_paragraph(self, style=None):
            if self.paragraph is None:
                self.paragraph = self.document.add_paragraph()
                if style:
                    try:
                        self.paragraph.style = style
                    except Exception:
                        pass
            return self.paragraph

        def finish(self):
            self.paragraph = None

        def handle_starttag(self, tag, attrs):
            tag = tag.lower()
            if tag in {"p", "div"}:
                self.finish()
                self.ensure_paragraph()
            elif tag == "blockquote":
                self.finish()
                self.ensure_paragraph("Intense Quote")
            elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                self.finish()
                self.ensure_paragraph("Heading " + tag[1])
            elif tag == "li":
                self.finish()
                self.ensure_paragraph("List Bullet")
            elif tag in {"strong", "b"}:
                self.bold = True
            elif tag in {"em", "i"}:
                self.italic = True
            elif tag == "u":
                self.underline = True
            elif tag in {"s", "strike"}:
                self.strike = True

        def handle_endtag(self, tag):
            tag = tag.lower()
            if tag in {"p", "div", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6", "li"}:
                self.finish()
            elif tag in {"strong", "b"}:
                self.bold = False
            elif tag in {"em", "i"}:
                self.italic = False
            elif tag == "u":
                self.underline = False
            elif tag in {"s", "strike"}:
                self.strike = False

        def handle_startendtag(self, tag, attrs):
            if tag.lower() == "br":
                self.ensure_paragraph().add_run().add_break()
            elif tag.lower() == "hr":
                p = self.ensure_paragraph()
                p.paragraph_format.space_after = 0
                p.add_run("────────────────────────").italic = True
                self.finish()

        def handle_data(self, data):
            if not data:
                return
            run = self.ensure_paragraph().add_run(data)
            run.bold = self.bold
            run.italic = self.italic
            run.underline = self.underline
            run.font.strike = self.strike

    builder = DocumentBuilder()
    builder.feed(content or "")
    builder.close()
    if not builder.document.paragraphs:
        builder.document.add_paragraph("")
    output = io.BytesIO()
    builder.document.save(output)
    output.seek(0)
    return output


@app.route("/api/projects/<int:project_id>/export/<fmt>", methods=["GET"])
@api_login_required
def export_project(project_id, fmt):
    user = current_user()
    db = get_db()
    try:
        project = project_for_user(db, project_id, user.id)
        if not project:
            return jsonify({"error": "Project not found."}), 404
        safe_title = re.sub(r"[^A-Za-z0-9._-]+", "_", project.title).strip("_") or "manuscript"
        if fmt == "txt":
            return send_file(
                io.BytesIO(html_to_text(project.content).encode("utf-8")),
                mimetype="text/plain; charset=utf-8",
                as_attachment=True,
                download_name=f"{safe_title}.txt",
            )
        if fmt == "docx":
            return send_file(
                html_to_docx(project.content),
                mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                as_attachment=True,
                download_name=f"{safe_title}.docx",
            )
        if fmt == "html":
            document = (
                "<!doctype html><html><head><meta charset='utf-8'><title>"
                + escape(project.title)
                + "</title><style>body{max-width:860px;margin:50px auto;padding:0 30px;background:#f7f5f2;color:#171411;font:17px/1.8 Georgia,serif}"
                "h1,h2,h3{color:#3f3027}hr{border:0;border-top:1px solid #c0ab9a;margin:2em 0}</style></head><body>"
                + project.content
                + "</body></html>"
            )
            return app.response_class(
                document,
                mimetype="text/html",
                headers={"Content-Disposition": f'attachment; filename="{safe_title}.html"'},
            )
        return jsonify({"error": "Unknown export format."}), 400
    finally:
        db.close()


@app.route("/analyze", methods=["POST"])
@api_login_required
def legacy_analyse():
    if "file" in request.files and request.files["file"].filename:
        try:
            _, content = parse_uploaded_file(request.files["file"])
            value = html_to_text(content)
        except Exception as exc:
            return jsonify({"error": f"Failed to read file: {exc}"}), 500
    else:
        value = request.form.get("text", "")
    return jsonify(analyse_text(value))


migrate_schema()

if __name__ == "__main__":
    app.run(debug=True)
